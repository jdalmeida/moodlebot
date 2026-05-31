"""Leitura de conteúdo do Moodle — cursos, atividades e arquivos.

Complementa o `scraper.py` (que só lê a timeline do dashboard). Aqui ficam as
funções que o agente (`src/agent/`) usa como ferramentas para entender o
contexto de uma tarefa: abrir a página da atividade, ler o enunciado completo,
listar materiais anexos e baixar/extrair o texto desses materiais.

Estratégia (igual ao resto do projeto): nada de Web Services REST (token
indisponível para usuário SSO). Tudo via a `MoodleSession` Playwright já
autenticada — scraping HTML das páginas e download via `context.request`
(os cookies de sessão carregam, inclusive em `pluginfile.php`).

Princípio defensivo herdado do scraper: vários seletores com fallback e
`logger.warning` em vez de explodir — o tema do UNISC pode divergir do Moodle
padrão. Ajustar seletores com `scripts/inspect_task.py` capturando o HTML real.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup, Tag
from loguru import logger
from pydantic import BaseModel

from ..config import settings
from ..models import TipoTarefa
from .scraper import _classify_tipo, _extract_cmid, _parse_prazo_ptbr
from .session import MoodleSession, SessionExpiredError


# Inverso de _MOD_NAME_TO_TIPO (scraper) — para montar /mod/<nome>/view.php.
_TIPO_TO_MOD: dict[TipoTarefa, str] = {
    TipoTarefa.ASSIGNMENT: "assign",
    TipoTarefa.QUIZ: "quiz",
    TipoTarefa.FORUM: "forum",
    TipoTarefa.CHOICE: "choice",
    TipoTarefa.WORKSHOP: "workshop",
}


# --------------------------------------------------------------------------- #
# Modelos
# --------------------------------------------------------------------------- #


class CursoMoodle(BaseModel):
    id: str
    nome: str
    url: str


class AtividadeResumo(BaseModel):
    cmid: str
    tipo: TipoTarefa
    nome: str
    url: str
    secao: str | None = None


class SecaoCurso(BaseModel):
    titulo: str
    resumo_texto: str | None = None
    atividades: list[AtividadeResumo] = []


class ArquivoAnexo(BaseModel):
    nome: str
    url: str
    mime_hint: str | None = None


class AtividadeDetalhe(BaseModel):
    cmid: str | None = None
    tipo: TipoTarefa = TipoTarefa.OTHER
    nome: str
    url: str
    enunciado_texto: str | None = None
    prazo: datetime | None = None
    status_envio: str | None = None
    anexos: list[ArquivoAnexo] = []
    links: list[str] = []


class ConteudoArquivo(BaseModel):
    url: str
    nome: str
    texto: str
    truncado: bool = False
    n_chars: int = 0


# --------------------------------------------------------------------------- #
# Helpers de parsing
# --------------------------------------------------------------------------- #


def _html_para_texto(node: Tag | None) -> str:
    """Texto limpo de um nó BeautifulSoup, colapsando espaços/linhas."""
    if node is None:
        return ""
    # Remove scripts/styles que sujam o get_text.
    for tag in node.find_all(["script", "style"]):
        tag.decompose()
    texto = node.get_text("\n", strip=True)
    # Colapsa runs de linhas em branco e espaços horizontais excessivos.
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


def _absolutizar(href: str, base_url: str) -> str:
    return href if href.startswith("http") else urljoin(base_url, href)


def _eh_link_interno(url: str, base_url: str) -> bool:
    try:
        return urlparse(url).netloc == urlparse(base_url).netloc
    except Exception:  # noqa: BLE001
        return False


def _montar_url_atividade(url_ou_cmid: str, tipo: TipoTarefa | None) -> str:
    """Aceita uma URL completa ou um cmid puro (+ tipo) e devolve a URL."""
    s = url_ou_cmid.strip()
    if s.startswith("http"):
        return s
    if s.isdigit():
        mod = _TIPO_TO_MOD.get(tipo or TipoTarefa.OTHER)
        if mod is None:
            raise ValueError(
                f"cmid {s!r} sem tipo mapeável — passe a URL completa da atividade."
            )
        base = settings.moodle_base_url.rstrip("/")
        return f"{base}/mod/{mod}/view.php?id={s}"
    # Caminho relativo do Moodle.
    return _absolutizar(s, settings.moodle_base_url)


async def _checar_login(page_url: str) -> None:
    if "login" in page_url.lower():
        raise SessionExpiredError(
            "Redirecionado para a página de login — sessão Moodle expirou. "
            "Rode `python scripts/first_login.py`."
        )


# --------------------------------------------------------------------------- #
# Cursos
# --------------------------------------------------------------------------- #


async def listar_cursos(session: MoodleSession) -> list[CursoMoodle]:
    """Lista os cursos em que o usuário está matriculado.

    Faz scraping do dashboard (`/my/`): o bloco "Visão geral dos cursos"
    renderiza cards com links `/course/view.php?id=N`. Dedupa por id.
    """
    page = await session.new_page()
    try:
        await page.goto(settings.moodle_dashboard_url, wait_until="domcontentloaded")
        await _checar_login(page.url)
        # O bloco de cursos vem por AJAX — espera um link de curso aparecer.
        try:
            await page.wait_for_selector(
                "a[href*='/course/view.php?id=']", timeout=8_000
            )
        except Exception:  # noqa: BLE001
            logger.info("Nenhum card de curso encontrado no dashboard.")
        html = await page.content()
    finally:
        await page.close()

    soup = BeautifulSoup(html, "lxml")
    vistos: dict[str, CursoMoodle] = {}
    for a in soup.select("a[href*='/course/view.php']"):
        href = a.get("href") or ""
        url = _absolutizar(href, settings.moodle_base_url)
        cid = _extract_cmid(url)
        if not cid or cid in vistos:
            continue
        nome = a.get_text(strip=True) or a.get("title") or f"Curso {cid}"
        vistos[cid] = CursoMoodle(id=cid, nome=nome, url=url)
    return list(vistos.values())


async def obter_curso(session: MoodleSession, curso_id: str) -> list[SecaoCurso]:
    """Lê a página do curso e devolve seções com suas atividades."""
    base = settings.moodle_base_url.rstrip("/")
    url = f"{base}/course/view.php?id={curso_id}"
    page = await session.new_page()
    try:
        await page.goto(url, wait_until="domcontentloaded")
        await _checar_login(page.url)
        html = await page.content()
    finally:
        await page.close()

    soup = BeautifulSoup(html, "lxml")
    secoes: list[SecaoCurso] = []

    section_nodes = soup.select("li.section, li[id^='section-'], .course-section")
    for sec in section_nodes:
        titulo_tag = sec.find(class_=re.compile(r"sectionname|section-title"))
        titulo = titulo_tag.get_text(strip=True) if titulo_tag else "(sem título)"

        resumo_tag = sec.find(class_="summary")
        resumo = _html_para_texto(resumo_tag) or None

        atividades: list[AtividadeResumo] = []
        for act in sec.select("li.activity, .activity"):
            link = act.find("a", href=re.compile(r"/mod/[a-z0-9]+/view\.php"))
            if link is None:
                continue
            href = link.get("href") or ""
            act_url = _absolutizar(href, settings.moodle_base_url)
            cmid = _extract_cmid(act_url)
            if not cmid:
                continue
            nome_tag = act.find(class_="instancename")
            nome = (
                nome_tag.get_text(strip=True)
                if nome_tag
                else link.get_text(strip=True)
            ) or "(sem nome)"
            atividades.append(
                AtividadeResumo(
                    cmid=cmid,
                    tipo=_classify_tipo(act_url),
                    nome=nome,
                    url=act_url,
                    secao=titulo,
                )
            )

        # Só registra seções que tenham conteúdo útil.
        if atividades or resumo:
            secoes.append(
                SecaoCurso(titulo=titulo, resumo_texto=resumo, atividades=atividades)
            )

    return secoes


# --------------------------------------------------------------------------- #
# Atividade (detalhe)
# --------------------------------------------------------------------------- #


# Seletores tentados em ordem para o enunciado/intro de uma atividade.
_SELETORES_ENUNCIADO = [
    "div.activity-description",
    "div[id^='intro']",
    "#intro",
    ".box.generalbox.boxaligncenter",
    ".box.py-3.generalbox",
    ".no-overflow",
]

# Reusa o seletor de status do submitter (mantém uma única fonte de verdade
# implícita: se mudar lá, ajustar aqui também).
_SELETOR_STATUS = "table.submissionstatustable"


def _extrair_enunciado(soup: BeautifulSoup) -> str:
    for sel in _SELETORES_ENUNCIADO:
        node = soup.select_one(sel)
        texto = _html_para_texto(node)
        if texto:
            return texto
    # Último recurso: o conteúdo principal da página, sem navegação.
    main = soup.select_one("div[role='main'], #region-main, section#region-main")
    return _html_para_texto(main)


def _extrair_prazo(soup: BeautifulSoup) -> datetime | None:
    """Procura uma data de entrega/encerramento no texto da página."""
    # Moodle pt-BR costuma rotular como "Data de entrega" / "Encerramento" /
    # "Disponível até". Varremos linhas que contenham esses rótulos.
    texto = soup.get_text("\n", strip=True)
    for linha in texto.splitlines():
        baixo = linha.lower()
        if any(
            rot in baixo
            for rot in ("data de entrega", "encerramento", "disponível até", "vencimento")
        ):
            prazo = _parse_prazo_ptbr(linha)
            if prazo is not None:
                return prazo
    return None


def _extrair_status(soup: BeautifulSoup) -> str | None:
    tabela = soup.select_one(_SELETOR_STATUS)
    if tabela is None:
        return None
    # A coluna de status costuma ser a última célula da primeira linha de dados.
    celulas = [c.get_text(strip=True) for c in tabela.select("td")]
    celulas = [c for c in celulas if c]
    if not celulas:
        return None
    # Junta as primeiras células dando uma visão do status sem inflar.
    return " | ".join(celulas[:4])


def _extrair_anexos_e_links(
    soup: BeautifulSoup, base_url: str
) -> tuple[list[ArquivoAnexo], list[str]]:
    main = soup.select_one("div[role='main'], #region-main") or soup
    anexos: dict[str, ArquivoAnexo] = {}
    links: list[str] = []
    vistos_links: set[str] = set()
    for a in main.find_all("a", href=True):
        href = a["href"]
        url = _absolutizar(href, base_url)
        if "pluginfile.php" in url:
            if url not in anexos:
                nome = a.get_text(strip=True) or url.rsplit("/", 1)[-1]
                anexos[url] = ArquivoAnexo(nome=nome, url=url)
        elif url.startswith("http") and not _eh_link_interno(url, base_url):
            if url not in vistos_links:
                vistos_links.add(url)
                links.append(url)
    return list(anexos.values()), links


def _nome_atividade(soup: BeautifulSoup) -> str:
    for sel in (".page-header-headings h1", "h1", "h2.main", "h2"):
        node = soup.select_one(sel)
        if node and node.get_text(strip=True):
            return node.get_text(strip=True)
    if soup.title and soup.title.string:
        return soup.title.string.strip()
    return "(atividade)"


async def obter_atividade(
    session: MoodleSession,
    url_ou_cmid: str,
    tipo: TipoTarefa | None = None,
) -> AtividadeDetalhe:
    """Abre a página de uma atividade e extrai enunciado, prazo, status,
    materiais anexos e links externos."""
    url = _montar_url_atividade(url_ou_cmid, tipo)
    page = await session.new_page()
    try:
        await page.goto(url, wait_until="domcontentloaded")
        await _checar_login(page.url)
        url_final = page.url
        html = await page.content()
    finally:
        await page.close()

    soup = BeautifulSoup(html, "lxml")
    anexos, links = _extrair_anexos_e_links(soup, settings.moodle_base_url)

    return AtividadeDetalhe(
        cmid=_extract_cmid(url_final),
        tipo=tipo or _classify_tipo(url_final),
        nome=_nome_atividade(soup),
        url=url_final,
        enunciado_texto=_extrair_enunciado(soup) or None,
        prazo=_extrair_prazo(soup),
        status_envio=_extrair_status(soup),
        anexos=anexos,
        links=links,
    )


# --------------------------------------------------------------------------- #
# Download + extração de texto
# --------------------------------------------------------------------------- #


def _detectar_tipo(nome: str, content_type: str) -> str:
    """Devolve uma chave canônica: pdf | docx | pptx | html | text | desconhecido."""
    ct = (content_type or "").lower()
    nome_baixo = nome.lower()
    if "pdf" in ct or nome_baixo.endswith(".pdf"):
        return "pdf"
    if "wordprocessingml" in ct or nome_baixo.endswith(".docx"):
        return "docx"
    if "presentationml" in ct or nome_baixo.endswith(".pptx"):
        return "pptx"
    if "html" in ct or nome_baixo.endswith((".html", ".htm")):
        return "html"
    if ct.startswith("text/") or nome_baixo.endswith((".txt", ".md", ".csv")):
        return "text"
    return "desconhecido"


def _extrair_texto(body: bytes, nome: str, content_type: str) -> str:
    """Extrai texto do arquivo. Imports são lazy para não pesar no import do
    módulo e para degradar com graça se um parser não estiver instalado."""
    kind = _detectar_tipo(nome, content_type)
    buf = io.BytesIO(body)
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(buf)
        partes = []
        for pagina in reader.pages:
            try:
                partes.append(pagina.extract_text() or "")
            except Exception:  # noqa: BLE001
                continue
        return "\n".join(partes).strip()
    if kind == "docx":
        import docx

        documento = docx.Document(buf)
        return "\n".join(p.text for p in documento.paragraphs).strip()
    if kind == "pptx":
        from pptx import Presentation

        prs = Presentation(buf)
        partes = []
        for i, slide in enumerate(prs.slides, start=1):
            partes.append(f"[slide {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    partes.append(shape.text_frame.text)
        return "\n".join(partes).strip()
    if kind == "html":
        return _html_para_texto(BeautifulSoup(body, "lxml"))
    if kind == "text":
        return body.decode("utf-8", errors="replace").strip()
    raise ValueError(f"tipo de arquivo não suportado para extração: {nome!r}")


async def baixar_arquivo(
    session: MoodleSession,
    url: str,
    max_bytes: int | None = None,
) -> ConteudoArquivo:
    """Baixa um arquivo do Moodle (autenticado) e extrai seu texto.

    Trunca em `settings.reader_max_text_chars`. Levanta `SessionExpiredError`
    se o download cair na página de login; `ValueError` para tipos não
    suportados ou arquivos grandes demais.
    """
    max_bytes = max_bytes or settings.reader_max_download_bytes
    req = session.request_context()
    resp = await req.get(url)

    if not resp.ok:
        if "login" in resp.url.lower():
            raise SessionExpiredError(
                "Download redirecionado para login — sessão expirou."
            )
        raise ValueError(f"download falhou ({resp.status}) para {url}")

    headers = resp.headers
    tamanho = headers.get("content-length")
    if tamanho is not None and tamanho.isdigit() and int(tamanho) > max_bytes:
        raise ValueError(
            f"arquivo grande demais ({tamanho} bytes > limite {max_bytes})"
        )

    body = await resp.body()
    if len(body) > max_bytes:
        raise ValueError(
            f"arquivo grande demais ({len(body)} bytes > limite {max_bytes})"
        )

    nome = url.rstrip("/").rsplit("/", 1)[-1].split("?")[0] or "arquivo"
    content_type = headers.get("content-type", "")
    texto = _extrair_texto(body, nome, content_type)

    limite = settings.reader_max_text_chars
    truncado = len(texto) > limite
    if truncado:
        texto = texto[:limite]

    return ConteudoArquivo(
        url=url,
        nome=nome,
        texto=texto,
        truncado=truncado,
        n_chars=len(texto),
    )
